"""发货单只导入查询所需的手机号及运单号，不保存姓名或地址。"""
from datetime import timedelta, timezone
from io import BytesIO
import json
import re
from zipfile import ZipFile

from fastapi import APIRouter, Depends, HTTPException, Request
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.dialects.mysql import insert
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from .db import get_db
from .models import Shipment, ShipmentImport, now
from .security import session_auth

router = APIRouter(prefix='/api/merchant/shipments')
MAX_BYTES, MAX_ROWS = 5 * 1024 * 1024, 5000
PHONE = re.compile(r'1[3-9]\d{9}', re.ASCII)


def scalar_text(value):
    if isinstance(value, bool) or value is None:
        return ''
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            return ''
        return str(int(value))
    return str(value).strip().lstrip("'")


def parse_excel(content):
    if len(content) > MAX_BYTES:
        raise HTTPException(413, '文件不能超过5MB。')
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(e.file_size for e in entries) > 30 * 1024 * 1024:
                raise HTTPException(400, '文件内容过大，请只保留发货数据后重新上传。')
            if any('vbaproject' in e.filename.lower() for e in entries):
                raise HTTPException(400, '只支持不含宏的xlsx文件。')
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, '无法读取文件，请上传有效的xlsx发货单。') from None
    records, issues, seen, sheets = [], [], set(), []
    visited = 0
    try:
        for sheet in workbook:
            # The provided workbook has formatting to XFD; bound columns, not max_column.
            rows = sheet.iter_rows(max_col=64, values_only=True)
            header = next(rows, ())
            labels = [scalar_text(x) for x in header]
            if '快递单号' not in labels:
                continue
            if labels.count('快递单号') != 1:
                raise HTTPException(400, '快递单号列重复，请整理表头。')
            tracking_col = labels.index('快递单号')
            candidates = [labels.index(name) for name in ('手机', '手机号', '固话') if name in labels]
            if not candidates:
                raise HTTPException(400, '缺少手机、手机号或固话列。')
            sheets.append(sheet.title)
            for row_number, row in enumerate(rows, 2):
                visited += 1
                if visited > MAX_ROWS:
                    raise HTTPException(400, '每次最多读取5000行，请拆分文件。')
                if all(x is None or x == '' for x in row):
                    continue
                tracking = scalar_text(row[tracking_col])
                numbers = {scalar_text(row[i]).replace(' ', '').replace('-', '')
                           for i in candidates if scalar_text(row[i])}
                phones = {p for p in numbers if PHONE.fullmatch(p)}
                reason = None
                if not re.fullmatch(r'[A-Za-z0-9-]{5,64}', tracking):
                    reason = '快递单号缺失或格式无效'
                elif isinstance(row[tracking_col], (int, float)) and len(tracking) > 15:
                    reason = '长快递单号为数字单元格，请改为文本并核对原始单号'
                elif len(phones) != 1:
                    reason = '手机号缺失、格式无效或多列号码冲突'
                elif tracking in seen:
                    reason = '快递单号重复'
                if reason:
                    if len(issues) < 50:
                        issues.append({'sheet': sheet.title, 'row': row_number, 'reason': reason})
                    continue
                seen.add(tracking)
                records.append({'tracking_number': tracking, 'phone': phones.pop()})
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, '表格结构无法读取，请重新导出xlsx。') from None
    finally:
        workbook.close()
    if not sheets or not records and not issues:
        raise HTTPException(400, '未找到包含快递单号和手机号的有效发货数据。')
    if issues:
        # Entire file is rejected; never silently import a partial delivery list.
        raise HTTPException(400, {'message': '请修正以下问题后重新上传，本次没有导入任何记录。', 'issues': issues})
    return records, sheets


@router.get('')
def status(request: Request, db: Session = Depends(get_db)):
    session_auth(request, db, merchant=True)
    count, latest = db.execute(select(func.count(), func.max(Shipment.updated_at))).one()
    return {'count': count, 'updated_at': latest.isoformat() + 'Z' if latest else None}


@router.post('/preview')
async def preview(request: Request, db: Session = Depends(get_db)):
    identity = session_auth(request, db, merchant=True)
    content = bytearray()
    async for part in request.stream():
        content.extend(part)
        if len(content) > MAX_BYTES:
            raise HTTPException(413, '文件不能超过5MB。')
    from starlette.concurrency import run_in_threadpool
    records, sheets = await run_in_threadpool(parse_excel, bytes(content))
    keys = [r['tracking_number'] for r in records]
    existing = set(db.scalars(select(Shipment.tracking_number).where(Shipment.tracking_number.in_(keys))))
    # Expired/previously committed drafts never need to retain phone numbers.
    db.query(ShipmentImport).filter(ShipmentImport.expires_at < now()).delete()
    draft = ShipmentImport(username=identity.username, payload=json.dumps(records),
                           expires_at=now() + timedelta(minutes=30))
    db.add(draft)
    db.commit()
    return {'import_id': draft.id, 'total': len(records), 'new': len(keys) - len(existing),
            'updated': len(existing), 'sheets': sheets,
            'sample': [{'tracking_number': r['tracking_number'], 'phone': r['phone'][:3] + '****' + r['phone'][-4:]}
                       for r in records[:5]]}


class Confirm(BaseModel):
    import_id: str


@router.post('/confirm')
def confirm(data: Confirm, request: Request, db: Session = Depends(get_db)):
    identity = session_auth(request, db, merchant=True)
    draft = db.scalar(select(ShipmentImport).where(ShipmentImport.id == data.import_id,
        ShipmentImport.username == identity.username).with_for_update())
    if not draft:
        raise HTTPException(404, '导入预览不存在。')
    if draft.committed_at:
        return {'ok': True, 'updated_at': draft.committed_at.isoformat() + 'Z'}
    if draft.expires_at <= now():
        raise HTTPException(409, '导入预览已过期，请重新上传。')
    stamp = now().replace(microsecond=0)
    for row in json.loads(draft.payload):
        db.execute(insert(Shipment).values(**row, updated_at=stamp).on_duplicate_key_update(
            phone=row['phone'], updated_at=stamp))
    draft.committed_at, draft.payload = stamp, '[]'
    db.commit()
    return {'ok': True, 'updated_at': stamp.isoformat() + 'Z'}


def lookup_reply(db, content, group=False):
    # Avoid matching 11 digits inside a longer order/tracking number.
    phones = re.findall(r'(?<!\d)(1[3-9]\d{9})(?!\d)', content, flags=re.ASCII)
    intent = any(word in content for word in ('查发货', '查快递', '查物流', '查询发货', '查询快递', '查询物流', '查订单', '查询订单'))
    if not phones and not intent:
        return None
    if group:
        return '请私聊机器人发送收件手机号查询发货记录，不要在群里公开手机号或地址。'
    if len(set(phones)) != 1:
        return '请发送一个完整的11位收件手机号，我帮你查发货记录。也可以发送“转人工”。'
    phone = phones[0]
    rows = list(db.scalars(select(Shipment).where(Shipment.phone == phone)
        .order_by(Shipment.updated_at.desc(), Shipment.tracking_number).limit(11)))
    if not rows:
        return '这份发货表中没有匹配记录，请核对收件手机号，或发送“转人工”让商家核实。'
    if len(rows) > 10:
        return '这个手机号对应的发货记录较多，请发送“转人工”让商家逐条核实。'
    stamp = max(r.updated_at for r in rows).replace(tzinfo=timezone.utc)
    date = stamp.astimezone(timezone(timedelta(hours=8))).strftime('%Y-%m-%d %H:%M')
    items = '\n'.join(f'{i}. 快递单号：{r.tracking_number}' for i, r in enumerate(rows, 1))
    return f'查到 {len(rows)} 条发货记录：\n{items}\n表格记录更新时间：{date}。\n这是商家上传的发货单记录，实时运输状态请到快递平台查询。'
